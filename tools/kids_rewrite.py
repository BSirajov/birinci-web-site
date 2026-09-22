# -*- coding: utf-8 -*-
"""Produce simplified (ages 10-14) Discoveries articles for a given language.

Source files are never modified. Rewritten copies land in a sibling folder
under their original filenames.

Only prose is rewritten: the summary paragraph, the key-facts list and the body
of the six content sections. The category line, the numbered title, the period
metadata line, every heading and the whole references section are copied
verbatim, so citations, hyperlinks, list bullets and formatting survive.

  python tools/kids_rewrite.py export --lang az
  python tools/kids_rewrite.py build  --lang az
  python tools/kids_rewrite.py verify --lang az
  python tools/kids_rewrite.py audit  --lang az [--detail]
"""
from __future__ import annotations

import argparse
import copy
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass, field
from functools import lru_cache
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]


@dataclass(frozen=True)
class LanguageSpec:
    key: str
    src_rel: str
    out_rel: str
    work_rel: str
    facts_head: str
    refs_head: str
    # Some corpora use more than one wording for the same six sections.
    section_variants: tuple[tuple[str, ...], ...]
    meta_prefix: str
    # Order of the three fixed lead paragraphs before the summary.
    lead_order: tuple[str, str, str]
    # "latin" protects foreign-script names; "caps" protects capitalised names;
    # "cyr_caps" protects both, comparing Cyrillic names by stem; "en_caps"
    # protects capitalised names in a Latin-script corpus, where every ordinary
    # word is Latin too and so cannot be protected wholesale.
    protect: str
    avoid: dict[str, str]
    gloss_markers: tuple[str, ...]
    alien_script: str = ""
    long_sentence: int = 24
    long_paragraph: int = 120

    @property
    def src_dir(self) -> Path:
        return ROOT / self.src_rel

    @property
    def out_dir(self) -> Path:
        return ROOT / self.out_rel

    @property
    def export_dir(self) -> Path:
        return ROOT / self.work_rel / "export"

    @property
    def rewrite_dir(self) -> Path:
        return ROOT / self.work_rel / "rewrite"

    def fold(self, text: str) -> str:
        return az_lower(text) if self.key == "az" else text.lower()

    def sections_for(self, texts: list[str]) -> tuple[str, ...]:
        for variant in self.section_variants:
            if variant[0] in texts:
                return variant
        raise ValueError("no known section heading set found")


RU = LanguageSpec(
    key="ru",
    src_rel="ru/discovery-articles",
    out_rel="ru/discovery-articles/Age10-14-Cursor",
    work_rel="tools/_ru_work",
    facts_head="Ключевые факты",
    refs_head="Источники и литература",
    section_variants=(
        (
            "Что это такое",
            "Кто это открыл или изобрёл",
            "Когда и как это появилось",
            "Научное значение",
            "Влияние на последующее развитие",
            "Как это изменило жизнь людей",
        ),
    ),
    meta_prefix="Период:",
    lead_order=("category", "title", "meta"),
    protect="latin",
    alien_script="A-Za-z",
    avoid={
        "поначалу": "сначала",
        "зачастую": "часто",
        "посредством": "с помощью",
        "ныне": "сейчас",
        "весьма": "очень",
        "дабы": "чтобы",
        "сие": "это",
        "сей": "этот",
        "оных": "этих",
        "надлежит": "нужно",
        "именуется": "называется",
        "осуществляется": "происходит",
        "вследствие": "из-за",
        "надобно": "нужно",
        "ибо": "потому что",
        "посему": "поэтому",
        "нежели": "чем",
        "столь": "так",
        "ежели": "если",
        "вкупе": "вместе",
        "представляет собой": "это",
        "в силу того": "потому что",
        "в настоящее время": "сейчас",
        "вероятнее всего": "скорее всего",
    },
    gloss_markers=(" — это ", " — то есть ", "называют"),
)

AZ = LanguageSpec(
    key="az",
    src_rel="az/discovery-articles",
    out_rel="az/discovery-articles/Age10-14-Cursor",
    work_rel="tools/_az_work",
    facts_head="Əsas faktlar",
    refs_head="Mənbələr və istinadlar",
    section_variants=(
        (
            "Bu nədir və necə işləyir?",
            "Kim kəşf və ya ixtira edib?",
            "Nə vaxt və necə yaranıb?",
            "Elmi əhəmiyyəti",
            "Sonrakı inkişafa təsiri",
            "İnsan həyatını necə dəyişdi?",
        ),
        (
            "Bu nədir",
            "Onu kim kəşf və ya ixtira edib",
            "Nə vaxt və necə meydana gəlib",
            "Elmi baxımdan əhəmiyyəti",
            "Sonrakı inkişaflara necə təsir göstərib",
            "İnsan həyatını necə dəyişib",
        ),
    ),
    meta_prefix="Dövr:",
    lead_order=("title", "category", "meta"),
    protect="caps",
    alien_script="\u0400-\u04ff",
    avoid={
        "vasitəsilə": "ilə",
        "mövcuddur": "var",
        "həyata keçirilir": "olur",
        "əhəmiyyət kəsb edir": "vacibdir",
        "nəticə etibarilə": "nəticədə",
        "bilavasitə": "birbaşa",
        "sözügedən": "bu",
        "adıçəkilən": "bu",
        "müvafiq olaraq": "uyğun olaraq",
        "dolayısıyla": "buna görə",
        "hasil olur": "əmələ gəlir",
        "vaqe olur": "baş verir",
        "qeyd etmək lazımdır ki": "",
        "yuxarıda qeyd edildiyi kimi": "",
    },
    gloss_markers=(
        " yəni ",
        "adlanır",
        "adlandırılır",
        "adlanan",
        "deyilir",
        "deyilən",
        "deməkdir",
        "termini",
    ),
)

KY = LanguageSpec(
    key="ky",
    src_rel="ky/discovery-articles",
    out_rel="ky/discovery-articles/Age10-14-Cursor",
    work_rel="tools/_ky_work",
    facts_head="Негизги фактылар",
    refs_head="Булактар жана шилтемелер",
    section_variants=(
        (
            "Бул эмне",
            "Аны ким ачкан же ойлоп тапкан",
            "Качан жана кантип пайда болгон",
            "Илимий жактан мааниси",
            "Кийинки өнүгүүлөргө тийгизген таасири",
            "Адамдардын жашоосун кантип өзгөрткөн",
        ),
    ),
    meta_prefix="Мезгил:",
    lead_order=("title", "category", "meta"),
    protect="cyr_caps",
    alien_script="A-Za-z",
    avoid={
        "болуп саналат": "болот",
        "жүзөгө ашырылат": "аткарылат",
        "ишке ашырылат": "болот",
        "мааниге ээ": "маанилүү",
        "тарабынан": "актив сүйлөм түзүү",
        "тиешелүү түрдө": "ошого жараша",
        "жогоруда белгиленгендей": "",
        "белгилей кетүү керек": "",
        "ушул тапта": "азыр",
        "аталмыш": "деп аталган",
        "сөз болуп жаткан": "бул",
    },
    gloss_markers=(
        " — бул ",
        " деген — ",
        "деп аталат",
        "деп аталган",
        "дегенди билдирет",
        "башкача айтканда",
        "тактап айтканда",
    ),
)

EN = LanguageSpec(
    key="en",
    src_rel="en/discovery-articles",
    out_rel="en/discovery-articles/Age10-14-Cursor",
    work_rel="tools/_en_work",
    facts_head="Key facts",
    refs_head="Sources and references",
    section_variants=(
        (
            "What it is",
            "Who discovered or invented it",
            "When and how it emerged",
            "Why it mattered scientifically",
            "How it shaped later developments",
            "How it changed human life",
        ),
    ),
    meta_prefix="Period:",
    lead_order=("title", "category", "meta"),
    protect="en_caps",
    alien_script="\u0400-\u04ff",
    avoid={
        "whilst": "while",
        "amongst": "among",
        "thus": "so",
        "hence": "so",
        "thereby": "so",
        "therein": "in it",
        "thereof": "of it",
        "herein": "here",
        "hitherto": "until now",
        "heretofore": "until now",
        "henceforth": "from now on",
        "albeit": "although",
        "notwithstanding": "despite",
        "aforementioned": "this",
        "utilise": "use",
        "utilize": "use",
        "commence": "start",
        "endeavour": "try",
        "ascertain": "find out",
        "elucidate": "explain",
        "disseminate": "spread",
        "myriad": "many",
        "plethora": "a lot of",
        "facilitate": "help",
        "necessitate": "need",
        "in order to": "to",
        "prior to": "before",
        "subsequent to": "after",
        "in the event that": "if",
        "due to the fact that": "because",
        "owing to the fact that": "because",
        "for the purpose of": "to",
        "by means of": "with",
        "with regard to": "about",
        "in the vicinity of": "near",
        "at the present time": "now",
        "it should be noted that": "",
        "as mentioned above": "",
    },
    gloss_markers=(
        " is a ",
        " is the ",
        " means ",
        "called",
        "known as",
        "in other words",
        "that is,",
        "for example",
    ),
)

SPECS = {spec.key: spec for spec in (RU, AZ, KY, EN)}

NUM_RE = re.compile(r"\d+")
LATIN_RE = re.compile(r"[A-Za-z][A-Za-z’'\-]+")
WORD_RE = re.compile(r"[^\W\d_]+(?:['’\-][^\W\d_]+)*", re.UNICODE)
SENT_SPLIT_RE = re.compile(r"[.!?…]+\s+|[.!?…]+$")
SENT_RE = re.compile(r"[^.!?…]+[.!?…]*")


def az_lower(text: str) -> str:
    """Lowercase using Azerbaijani rules, avoiding Python's dotted-I quirk."""
    return text.replace("İ", "i").replace("I", "ı").lower()


def az_stem(word: str) -> str:
    """Root-ish stem of an Azerbaijani word.

    Azerbaijani glues case endings onto names, including onto abbreviations
    across a hyphen (`DNT-ni`, `Yerdə`). Cutting at the hyphen and keeping only
    the first few letters lets a name match whatever ending it later takes.
    """
    return re.split(r"[-’']", az_lower(word), maxsplit=1)[0][:3]


CYRILLIC_RE = re.compile(r"[\u0400-\u04ff]")
EN_ENDINGS = ("ian", "ish", "ese", "ic", "an", "s")


def en_stem(word: str) -> str:
    """Root-ish stem of an English word.

    English names take no case endings, but they do spawn adjectives, and a
    rewrite may fairly swap one form for the other (`Egyptian scribes` becoming
    `scribes in Egypt`). Dropping a nationality ending and keeping the opening
    letters lets the two forms match.
    """
    base = re.split(r"[’']", word.lower(), maxsplit=1)[0]
    base = re.split(r"-", base, maxsplit=1)[0]
    for ending in EN_ENDINGS:
        if base.endswith(ending) and len(base) - len(ending) >= 4:
            base = base[: -len(ending)]
            break
    return base[:4]


def ky_stem(word: str) -> str:
    """Root-ish stem of a Kyrgyz word.

    Kyrgyz glues case endings onto names just as Azerbaijani does (`Фарадей` ->
    `Фарадейдин`), so a name is matched on its opening letters. The stem must be
    no longer than the shortest name in the corpus (`Рим`), otherwise a suffixed
    form and its bare form stem differently and stop matching each other.
    """
    return re.split(r"[-’']", word.lower(), maxsplit=1)[0][:3]


# --------------------------------------------------------------------------- #
# reading source documents
# --------------------------------------------------------------------------- #


def para_style(para) -> str:
    return para.style.name if para.style else "Normal"


def visible_paragraphs(doc) -> list:
    return [p for p in doc.paragraphs if (p.text or "").strip()]


def article_blocks(doc, spec: LanguageSpec) -> dict:
    paras = visible_paragraphs(doc)
    texts = [(p.text or "").strip() for p in paras]
    heads = spec.sections_for(texts)

    missing = [h for h in (spec.facts_head, spec.refs_head, *heads) if h not in texts]
    if missing:
        raise ValueError(f"missing headings: {missing}")

    i_facts = texts.index(spec.facts_head)
    i_refs = texts.index(spec.refs_head)
    if i_facts < 4:
        raise ValueError("unexpected lead block")

    lead = dict(zip(spec.lead_order, paras[:3]))
    if not texts[2].startswith(spec.meta_prefix):
        raise ValueError(f"unexpected metadata line: {texts[2]!r}")

    idx = {h: texts.index(h) for h in heads}
    sections = []
    for n, head in enumerate(heads):
        end = idx[heads[n + 1]] if n + 1 < len(heads) else i_refs
        sections.append({"heading": head, "paragraphs": paras[idx[head] + 1 : end]})

    return {
        "category": lead["category"],
        "title": lead["title"],
        "meta": lead["meta"],
        "summary": paras[3:i_facts],
        "facts": paras[i_facts + 1 : idx[heads[0]]],
        "sections": sections,
        "refs": paras[i_refs:],
        "headings": heads,
    }


def export_one(path: Path, spec: LanguageSpec) -> dict:
    blocks = article_blocks(Document(str(path)), spec)
    return {
        "file": path.name,
        "category": blocks["category"].text.strip(),
        "title": blocks["title"].text.strip(),
        "meta": blocks["meta"].text.strip(),
        "summary": [p.text.strip() for p in blocks["summary"]],
        "facts": [p.text.strip() for p in blocks["facts"]],
        "sections": [
            {
                "heading": s["heading"],
                "paragraphs": [p.text.strip() for p in s["paragraphs"]],
            }
            for s in blocks["sections"]
        ],
        "references": [
            {"style": para_style(p), "text": p.text.strip()} for p in blocks["refs"]
        ],
    }


def cmd_export(spec: LanguageSpec) -> int:
    spec.export_dir.mkdir(parents=True, exist_ok=True)
    count = 0
    for path in sorted(spec.src_dir.glob("*.docx")):
        data = export_one(path, spec)
        dest = spec.export_dir / f"{path.stem}.json"
        dest.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        count += 1
    print(f"exported {count} articles to {spec.export_dir.relative_to(ROOT)}")
    return 0


# --------------------------------------------------------------------------- #
# writing rewritten documents
# --------------------------------------------------------------------------- #


def _clone_with_text(template, text: str):
    """Copy a paragraph's XML, keeping style and list numbering, with new text."""
    new_p = copy.deepcopy(template._p)
    for child in list(new_p):
        if child.tag in (
            qn("w:r"),
            qn("w:hyperlink"),
            qn("w:bookmarkStart"),
            qn("w:bookmarkEnd"),
        ):
            new_p.remove(child)
    run_template = next(template._p.iterchildren(qn("w:r")), None)
    if run_template is None:
        raise ValueError("template paragraph has no run")
    new_run = copy.deepcopy(run_template)
    for node in list(new_run.iterchildren(qn("w:t"))):
        new_run.remove(node)
    t = new_run.makeelement(qn("w:t"), {})
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    new_p.append(new_run)
    return new_p


def _replace_group(paragraphs: list, texts: list[str]) -> None:
    if not paragraphs:
        raise ValueError("cannot replace an empty paragraph group")
    if not texts:
        raise ValueError("refusing to write an empty paragraph group")
    template = paragraphs[0]
    for text in texts:
        template._p.addprevious(_clone_with_text(template, text))
    for para in paragraphs:
        para._p.getparent().remove(para._p)


def build_one(path: Path, rewrite: dict, spec: LanguageSpec) -> Path:
    doc = Document(str(path))
    blocks = article_blocks(doc, spec)
    heads = blocks["headings"]

    summary = [s.strip() for s in rewrite.get("summary", []) if s.strip()]
    facts = [s.strip() for s in rewrite.get("facts", []) if s.strip()]
    sections = rewrite.get("sections") or []
    if len(sections) != len(heads):
        raise ValueError(f"{path.name}: expected {len(heads)} sections, got {len(sections)}")
    for section, head in zip(sections, heads):
        given = section.get("heading", head)
        if given != head:
            raise ValueError(f"{path.name}: heading mismatch {given!r} != {head!r}")

    _replace_group(blocks["summary"], summary)
    _replace_group(blocks["facts"], facts)
    for source_section, new_section in zip(blocks["sections"], sections):
        paragraphs = [s.strip() for s in new_section.get("paragraphs", []) if s.strip()]
        _replace_group(source_section["paragraphs"], paragraphs)

    spec.out_dir.mkdir(parents=True, exist_ok=True)
    dest = spec.out_dir / path.name
    doc.save(str(dest))
    return dest


def cmd_build(spec: LanguageSpec) -> int:
    names = [f"{p.stem}.docx" for p in sorted(spec.rewrite_dir.glob("*.json"))]
    built = 0
    skipped = 0
    for name in names:
        src = spec.src_dir / name
        if not src.is_file():
            print(f"  MISSING SOURCE {name}")
            continue
        payload = spec.rewrite_dir / f"{Path(name).stem}.json"
        # Rewrite files may be written concurrently, and an output document may
        # be held open by Word; skip whatever cannot be read or written and keep
        # going, so one locked file cannot stall the whole corpus.
        try:
            rewrite = json.loads(payload.read_text(encoding="utf-8"))
            build_one(src, rewrite, spec)
        except (json.JSONDecodeError, ValueError, KeyError, OSError) as exc:
            skipped += 1
            print(f"  SKIP {name}: {exc}")
            continue
        built += 1
    print(f"built {built} articles into {spec.out_dir.relative_to(ROOT)}")
    if skipped:
        print(f"skipped {skipped} unreadable or incomplete rewrite files")
    return 0 if built else 1


# --------------------------------------------------------------------------- #
# checks
# --------------------------------------------------------------------------- #


def body_text(data: dict) -> str:
    paras = (
        list(data["summary"])
        + list(data["facts"])
        + [p for s in data["sections"] for p in s["paragraphs"]]
    )
    # Bullet items often carry no final period; without one the sentence
    # statistics would treat consecutive items as a single run-on sentence.
    closed = []
    for para in paras:
        para = para.strip()
        if para and para[-1] not in ".!?…":
            para += "."
        closed.append(para)
    return " ".join(closed)


def name_stem(word: str, spec: LanguageSpec) -> str:
    if spec.protect == "caps":
        return az_stem(word)
    return en_stem(word) if spec.protect == "en_caps" else ky_stem(word)


@lru_cache(maxsize=None)
def common_word_stems(key: str) -> frozenset[str]:
    """Stems the corpus shows to be ordinary words rather than proper names.

    Two signals, both drawn from all 121 articles at once. A word that ever
    appears lowercase (`бирок`) is ordinary. So is one that only ever opens
    sentences yet does so across many articles (`Ошентсе`, `Ошентип`), because a
    real name turns up mid-sentence somewhere, while a discourse connective
    never does.
    """
    spec = SPECS[key]
    lowercase: set[str] = set()
    mid_sentence: set[str] = set()
    openers: dict[str, set[str]] = {}
    for path in sorted(spec.src_dir.glob("*.docx")):
        text = body_text(export_one(path, spec))
        for sentence in SENT_RE.findall(text):
            for position, word in enumerate(WORD_RE.findall(sentence)):
                stem = name_stem(word, spec)
                if spec.fold(word[0]) == word[0]:
                    lowercase.add(stem)
                elif position:
                    mid_sentence.add(stem)
                else:
                    openers.setdefault(stem, set()).add(path.name)
    habitual = {s for s, files in openers.items() if len(files) >= 3}
    return frozenset(lowercase | (habitual - mid_sentence))


def proper_names(text: str, spec: LanguageSpec) -> list[str]:
    """Capitalised words that look like names rather than ordinary words.

    A word capitalised only because it opens a sentence is usually an ordinary
    word (`Бирок`, `Ошентсе`), and dropping such a connective is a fair edit. So
    a sentence-initial word counts only when it never appears lowercase in the
    corpus, which keeps names mentioned just once (`Израилдеги`) protected.

    A capital anywhere but the first letter marks a name whatever its position,
    which catches the forms that start lowercase (`mRNA`, `iPhone`).
    """
    common = common_word_stems(spec.key)
    names = []
    for sentence in SENT_RE.findall(text):
        for position, word in enumerate(WORD_RE.findall(sentence)):
            inner_capital = word[1:] != spec.fold(word[1:])
            if spec.fold(word[0]) == word[0]:
                if inner_capital:
                    names.append(word)
                continue
            if position or inner_capital or name_stem(word, spec) not in common:
                names.append(word)
    return names


def required_tokens(text: str, spec: LanguageSpec) -> set[str]:
    """Facts that must survive a rewrite: numbers and proper names."""
    tokens = {f"#{n}" for n in NUM_RE.findall(expand_decades(text))}
    if spec.protect in ("latin", "cyr_caps"):
        tokens |= {w for w in LATIN_RE.findall(text) if len(w) > 1}
    if spec.protect == "caps":
        # Azerbaijani prose is Latin script throughout, so names are spotted by
        # capitalisation and compared by stem, since suffixes attach to them.
        for word in proper_names(text, spec):
            if len(word) >= 3:
                tokens.add(az_stem(word))
    elif spec.protect == "cyr_caps":
        # Kyrgyz names are Cyrillic and take suffixes, so they are matched by
        # stem; Latin tokens above are already required verbatim.
        for word in proper_names(text, spec):
            if len(word) >= 3 and CYRILLIC_RE.match(word):
                tokens.add(f"~{ky_stem(word)}")
    elif spec.protect == "en_caps":
        for word in proper_names(text, spec):
            if len(word) >= 3:
                tokens.add(f"~{en_stem(word)}")
    return tokens


def available_tokens(text: str, spec: LanguageSpec) -> set[str]:
    """Everything the rewrite offers to satisfy the required tokens.

    Capitalisation cannot be required on the rewrite side: a word that opened a
    sentence in the source may legitimately sit mid-sentence afterwards.
    """
    tokens = {f"#{n}" for n in NUM_RE.findall(expand_decades(text))}
    if spec.protect in ("latin", "cyr_caps"):
        tokens |= {w for w in LATIN_RE.findall(text) if len(w) > 1}
    if spec.protect == "caps":
        tokens |= {az_stem(w) for w in WORD_RE.findall(text) if len(w) >= 3}
    elif spec.protect == "cyr_caps":
        tokens |= {f"~{ky_stem(w)}" for w in WORD_RE.findall(text) if len(w) >= 3}
    elif spec.protect == "en_caps":
        tokens |= {f"~{en_stem(w)}" for w in WORD_RE.findall(text) if len(w) >= 3}
    return tokens


DECADE_RANGE_RE = re.compile(r"\b(1\d)(\d)0s\s*(?:[\u2013-]|and|to)\s*(\d)0s\b")


def expand_decades(text: str) -> str:
    """Write abbreviated decade ranges out in full before counting numbers.

    Sources shorten ranges to `1980s-90s`, where the bare `90s` yields the
    number 90. Requiring a literal 90 would force a rewrite to keep the same
    shorthand, so both sides are expanded to `1980s 1990s` first and either
    phrasing then satisfies the check.
    """
    return DECADE_RANGE_RE.sub(lambda m: f"{m[1]}{m[2]}0s {m[1]}{m[3]}0s", text)


def avg_sentence_words(text: str) -> float:
    parts = [s for s in SENT_SPLIT_RE.split(text) if s and s.strip()]
    if not parts:
        return 0.0
    return sum(len(s.split()) for s in parts) / len(parts)


def verify_one(name: str, spec: LanguageSpec) -> dict:
    result: dict = {"file": name, "errors": [], "warnings": []}
    out_path = spec.out_dir / name
    if not out_path.is_file():
        result["errors"].append("output missing")
        return result
    a = export_one(spec.src_dir / name, spec)
    b = export_one(out_path, spec)

    if a["category"] != b["category"]:
        result["errors"].append("category line changed")
    if a["title"] != b["title"]:
        result["errors"].append("title changed")
    if a["meta"] != b["meta"]:
        result["errors"].append("metadata line changed")
    if [s["heading"] for s in a["sections"]] != [s["heading"] for s in b["sections"]]:
        result["errors"].append("section headings changed")
    if a["references"] != b["references"]:
        result["errors"].append("references changed")
    if len(b["facts"]) != len(a["facts"]):
        result["warnings"].append(f"key facts {len(a['facts'])} -> {len(b['facts'])}")

    src_text, out_text = body_text(a), body_text(b)
    if not out_text.strip():
        result["errors"].append("empty body")
        return result

    lost = sorted(required_tokens(src_text, spec) - available_tokens(out_text, spec))
    if lost:
        pretty = [t.lstrip("#~") for t in lost]
        result["errors"].append("lost names or numbers: " + ", ".join(pretty[:12]))
    if spec.alien_script:
        stray = re.findall(rf"[{spec.alien_script}]+", out_text)
        if stray and not re.findall(rf"[{spec.alien_script}]+", src_text):
            result["errors"].append("wrong-script text: " + ", ".join(sorted(set(stray))[:6]))

    src_avg, out_avg = avg_sentence_words(src_text), avg_sentence_words(out_text)
    result.update(
        src_avg_sentence=round(src_avg, 1),
        out_avg_sentence=round(out_avg, 1),
        src_words=len(src_text.split()),
        out_words=len(out_text.split()),
    )
    if out_avg > 14.0:
        result["warnings"].append(f"long sentences (avg {out_avg:.1f} words)")
    if out_avg > src_avg:
        result["warnings"].append("sentences not shorter than source")
    if result["out_words"] < 0.85 * result["src_words"]:
        result["warnings"].append("output much shorter than source")
    return result


def cmd_verify(spec: LanguageSpec) -> int:
    names = [p.name for p in sorted(spec.src_dir.glob("*.docx"))]
    done = [n for n in names if (spec.out_dir / n).is_file()]
    print(f"source articles: {len(names)}")
    print(f"rewritten:       {len(done)}")
    missing = [n for n in names if n not in done]
    if missing:
        print(f"not yet rewritten ({len(missing)}):")
        for name in missing:
            print(f"  {name}")
    errors = warned = 0
    src_avg: list[float] = []
    out_avg: list[float] = []
    for name in done:
        res = verify_one(name, spec)
        if res["errors"]:
            errors += 1
            print(f"ERROR {name}")
            for item in res["errors"]:
                print(f"    {item}")
        if res["warnings"]:
            warned += 1
            print(f"WARN  {name}: {'; '.join(res['warnings'])}")
        if "src_avg_sentence" in res:
            src_avg.append(res["src_avg_sentence"])
            out_avg.append(res["out_avg_sentence"])
    if src_avg:
        print(
            f"avg sentence length: source {sum(src_avg)/len(src_avg):.1f} -> "
            f"rewritten {sum(out_avg)/len(out_avg):.1f} words"
        )
    print(f"files with errors: {errors}; with warnings: {warned}")
    return 1 if errors or missing else 0


_WORD_CACHE: dict[str, re.Pattern] = {}


def _whole_word(phrase: str, haystack: str, spec: LanguageSpec) -> bool:
    key = f"{spec.key}:{phrase}"
    pattern = _WORD_CACHE.get(key)
    if pattern is None:
        letters = "а-яёА-ЯЁ" if spec.key == "ru" else r"^\W\d_"
        cls = f"[{letters}]" if spec.key == "ru" else r"[^\W\d_]"
        pattern = re.compile(rf"(?<!{cls}){re.escape(phrase)}(?!{cls})")
        _WORD_CACHE[key] = pattern
    return bool(pattern.search(haystack))


def _stems(text: str) -> set[str]:
    return {w[:7] for w in WORD_RE.findall(text.lower()) if len(w) >= 9}


def audit_one(name: str, spec: LanguageSpec, detail: bool) -> dict:
    a = export_one(spec.src_dir / name, spec)
    b = export_one(spec.out_dir / name, spec)
    out_paras = (
        list(b["summary"]) + list(b["facts"]) + [p for s in b["sections"] for p in s["paragraphs"]]
    )
    src_text, out_text = body_text(a), body_text(b)
    issues: list[str] = []

    lowered = spec.fold(out_text)
    found = [w for w in spec.avoid if _whole_word(w, lowered, spec)]
    if found:
        issues.append(
            "bookish wording: "
            + ", ".join(f"{w} → {spec.avoid[w] or 'drop'}" for w in found[:6])
        )

    long_sents = [s.strip() for s in SENT_RE.findall(out_text) if len(s.split()) > spec.long_sentence]
    if long_sents:
        issues.append(f"{len(long_sents)} sentence(s) over {spec.long_sentence} words")
        if detail:
            for s in long_sents[:3]:
                issues.append(f"    · {s[:150]}")

    long_paras = [p for p in out_paras if len(p.split()) > spec.long_paragraph]
    if long_paras:
        issues.append(f"{len(long_paras)} paragraph(s) over {spec.long_paragraph} words")

    dupes = [p for p, n in Counter(out_paras).items() if n > 1]
    if dupes:
        issues.append(f"{len(dupes)} duplicated paragraph(s)")

    thin = [
        s["heading"]
        for s in b["sections"]
        if len(s["paragraphs"]) < 2 and len(" ".join(s["paragraphs"]).split()) > 60
    ]
    if thin:
        issues.append("single dense paragraph in: " + ", ".join(thin))

    glosses = sum(lowered.count(m) for m in spec.gloss_markers)
    if glosses == 0:
        issues.append("no term explanations found")

    dropped = sorted(_stems(src_text) - _stems(out_text))
    coverage = 1 - len(dropped) / max(len(_stems(src_text)), 1)
    if coverage < 0.45:
        issues.append(f"only {coverage:.0%} of distinctive source words have a match")
    if detail and dropped:
        issues.append("    dropped word stems: " + ", ".join(dropped[:15]))

    src_sents = [s for s in SENT_RE.findall(src_text) if s.strip()]
    out_sents = [s for s in SENT_RE.findall(out_text) if s.strip()]
    return {
        "issues": issues,
        "src_words": len(src_text.split()),
        "out_words": len(out_text.split()),
        "src_sent": sum(len(s.split()) for s in src_sents) / max(len(src_sents), 1),
        "out_sent": sum(len(s.split()) for s in out_sents) / max(len(out_sents), 1),
        "glosses": glosses,
        "coverage": coverage,
    }


def cmd_audit(spec: LanguageSpec, detail: bool) -> int:
    names = [p.name for p in sorted(spec.src_dir.glob("*.docx")) if (spec.out_dir / p.name).is_file()]
    if not names:
        print("nothing to audit yet")
        return 1
    flagged = 0
    unreadable: list[str] = []
    totals = dict(src_words=0, out_words=0, src_sent=0.0, out_sent=0.0, glosses=0, coverage=0.0)
    for name in names:
        # A concurrent build may be mid-save on this file.
        try:
            res = audit_one(name, spec, detail)
        except Exception as exc:  # noqa: BLE001 - report and keep auditing
            unreadable.append(f"{name}: {type(exc).__name__}")
            continue
        for key in totals:
            totals[key] += res[key]
        if res["issues"]:
            flagged += 1
            print(f"\n{name}")
            for item in res["issues"]:
                print(f"  {item}")
    n = len(names) - len(unreadable)
    print("\n" + "=" * 70)
    if unreadable:
        print(f"could not read {len(unreadable)} file(s), rerun to include them:")
        for item in unreadable:
            print(f"  {item}")
    if not n:
        return 1
    print(f"audited {n} articles; {flagged} with notes")
    grow = 100 * (totals["out_words"] - totals["src_words"]) / max(totals["src_words"], 1)
    print(f"words: {totals['src_words']} -> {totals['out_words']} ({grow:+.0f}%)")
    print(f"avg sentence: {totals['src_sent']/n:.1f} -> {totals['out_sent']/n:.1f} words")
    print(f"term explanations: {totals['glosses']} total, {totals['glosses']/n:.1f} per article")
    print(f"distinctive source words matched: {totals['coverage']/n:.0%} average")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="cmd", required=True)
    for name in ("export", "build", "verify", "audit"):
        p = sub.add_parser(name)
        p.add_argument("--lang", required=True, choices=sorted(SPECS))
        if name == "audit":
            p.add_argument("--detail", action="store_true")
    args = parser.parse_args()
    spec = SPECS[args.lang]
    if args.cmd == "export":
        return cmd_export(spec)
    if args.cmd == "build":
        return cmd_build(spec)
    if args.cmd == "verify":
        return cmd_verify(spec)
    return cmd_audit(spec, args.detail)


if __name__ == "__main__":
    sys.exit(main())
